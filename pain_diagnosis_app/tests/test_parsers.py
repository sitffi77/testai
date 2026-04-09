"""Tests for document parsers module."""

import pytest
from pathlib import Path
import tempfile
import os

from nlp.models import DocumentText


class TestPDFParser:
    """Tests for PDF parser (when pdfplumber is available)."""
    
    def test_pdf_parser_import(self):
        """Test that PDFParser can be imported."""
        try:
            from parsers.pdf_parser import PDFParser
            assert PDFParser is not None
        except ImportError:
            pytest.skip("pdfplumber not installed")
    
    def test_pdf_parser_initialization(self):
        """Test PDFParser initializes correctly."""
        try:
            from parsers.pdf_parser import PDFParser
            parser = PDFParser()
            assert parser is not None
        except ImportError:
            pytest.skip("pdfplumber not installed")


class TestDOCXParser:
    """Tests for DOCX parser (when python-docx is available)."""
    
    def test_docx_parser_import(self):
        """Test that DOCXParser can be imported."""
        try:
            from parsers.docx_parser import DOCXParser
            assert DOCXParser is not None
        except ImportError:
            pytest.skip("python-docx not installed")
    
    def test_docx_parser_initialization(self):
        """Test DOCXParser initializes correctly."""
        try:
            from parsers.docx_parser import DOCXParser
            parser = DOCXParser()
            assert parser is not None
        except ImportError:
            pytest.skip("python-docx not installed")
    
    def test_docx_parse_simple_document(self):
        """Test parsing a simple DOCX document."""
        try:
            from parsers.docx_parser import DOCXParser
            from docx import Document
            
            # Create temporary DOCX file
            with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
                tmp_path = tmp.name
            
            try:
                # Create document with content
                doc = Document()
                doc.add_heading('Медицинская выписка', 0)
                doc.add_paragraph('Пациент: Иванов И.И., 45 лет.')
                doc.add_paragraph('Жалобы: Боль в спине.')
                doc.save(tmp_path)
                
                # Parse the document
                parser = DOCXParser()
                result = parser.parse(tmp_path)
                
                # Verify results
                assert isinstance(result, DocumentText)
                assert 'Иванов' in result.content
                assert '45 лет' in result.content
                assert result.age == 45
                
            finally:
                # Cleanup
                if os.path.exists(tmp_path):
                    os.unlink(tmp_path)
                    
        except ImportError:
            pytest.skip("python-docx not installed")
    
    def test_docx_parse_empty_document(self):
        """Test parsing an empty DOCX document."""
        try:
            from parsers.docx_parser import DOCXParser
            from docx import Document
            
            # Create empty temporary DOCX file
            with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
                tmp_path = tmp.name
            
            try:
                doc = Document()
                doc.save(tmp_path)
                
                parser = DOCXParser()
                result = parser.parse(tmp_path)
                
                assert isinstance(result, DocumentText)
                assert result.content == ""
                
            finally:
                if os.path.exists(tmp_path):
                    os.unlink(tmp_path)
                    
        except ImportError:
            pytest.skip("python-docx not installed")
    
    def test_docx_parse_with_table(self):
        """Test parsing DOCX with tables."""
        try:
            from parsers.docx_parser import DOCXParser
            from docx import Document
            
            with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
                tmp_path = tmp.name
            
            try:
                doc = Document()
                doc.add_paragraph('Анализы пациента:')
                
                table = doc.add_table(rows=2, cols=2)
                table.cell(0, 0).text = 'Показатель'
                table.cell(0, 1).text = 'Значение'
                table.cell(1, 0).text = 'Температура'
                table.cell(1, 1).text = '37.5°C'
                
                doc.save(tmp_path)
                
                parser = DOCXParser()
                result = parser.parse(tmp_path)
                
                assert 'Температура' in result.content
                assert '37.5' in result.content
                
            finally:
                if os.path.exists(tmp_path):
                    os.unlink(tmp_path)
                    
        except ImportError:
            pytest.skip("python-docx not installed")


class TestAgeExtraction:
    """Tests for age extraction from documents."""
    
    def test_age_extraction_patterns(self):
        """Test various age extraction patterns."""
        from nlp.context_analyzer import ContextAnalyzer
        
        analyzer = ContextAnalyzer([], [])
        
        test_cases = [
            ("Пациент 45 лет", 45),
            ("возраст 62 года", 62),
            ("30-летний мужчина", 30),
            ("пациент 18 лет", 18),
            ("женщина 55 лет", 55),
        ]
        
        for text, expected_age in test_cases:
            result = analyzer.extract_age(text)
            assert result == expected_age, f"Failed for: {text}"
    
    def test_no_age_extraction(self):
        """Test when no age is present."""
        from nlp.context_analyzer import ContextAnalyzer
        
        analyzer = ContextAnalyzer([], [])
        text = "Пациент жалуется на боль"
        
        result = analyzer.extract_age(text)
        assert result is None


class TestTemperatureExtraction:
    """Tests for temperature extraction from documents."""
    
    def test_temperature_extraction_patterns(self):
        """Test various temperature extraction patterns."""
        from nlp.context_analyzer import ContextAnalyzer
        
        analyzer = ContextAnalyzer([], [])
        
        test_cases = [
            ("Температура 37.5°C", 37.5),
            ("t=38.2C", 38.2),
            ("температура: 36.8°C", 36.8),
            ("субфебрилитет", 37.1),  # Default for subfebrile
            ("лихорадка", 37.1),  # Default for fever
        ]
        
        for text, expected_temp in test_cases:
            result = analyzer.extract_temperature(text)
            assert result == expected_temp, f"Failed for: {text}"
    
    def test_no_temperature_extraction(self):
        """Test when no temperature is present."""
        from nlp.context_analyzer import ContextAnalyzer
        
        analyzer = ContextAnalyzer([], [])
        text = "Пациент чувствует себя хорошо"
        
        result = analyzer.extract_temperature(text)
        assert result is None


class TestErrorHandling:
    """Tests for error handling in parsers."""
    
    def test_file_not_found(self):
        """Test handling of non-existent files."""
        try:
            from parsers.docx_parser import DOCXParser
            
            parser = DOCXParser()
            
            with pytest.raises(FileNotFoundError):
                parser.parse("/nonexistent/path/file.docx")
                
        except ImportError:
            pytest.skip("python-docx not installed")
    
    def test_invalid_extension(self):
        """Test handling of invalid file extensions."""
        try:
            from parsers.docx_parser import DOCXParser
            
            parser = DOCXParser()
            
            with pytest.raises(ValueError):
                parser.parse("/some/path/file.txt")
                
        except ImportError:
            pytest.skip("python-docx not installed")


if __name__ == '__main__':
    pytest.main([__file__, '-v'])
