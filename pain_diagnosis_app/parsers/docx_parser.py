"""DOCX document parser using python-docx."""

import logging
from pathlib import Path
from typing import List, Optional

try:
    from docx import Document
except ImportError:
    Document = None
    logging.warning("python-docx not installed. DOCX parsing will be unavailable.")

from nlp.models import DocumentText

logger = logging.getLogger(__name__)


class DOCXParser:
    """Парсер DOCX документов.
    
    Использует python-docx для извлечения текста из DOCX файлов.
    Поддерживает извлечение метаданных (возраст, температура) из текста.
    """
    
    def __init__(self):
        """Инициализация DOCX парсера."""
        if Document is None:
            raise ImportError(
                "python-docx не установлен. Установите: pip install python-docx"
            )
        
        logger.info("DOCXParser инициализирован")
    
    def parse(self, file_path: str) -> DocumentText:
        """Извлечение текста из DOCX файла.
        
        Args:
            file_path: Путь к DOCX файлу.
            
        Returns:
            DocumentText с извлечённым текстом и метаданными.
            
        Raises:
            FileNotFoundError: Если файл не найден.
            ValueError: Если файл не является валидным DOCX.
        """
        path = Path(file_path)
        
        if not path.exists():
            logger.error(f"Файл не найден: {file_path}")
            raise FileNotFoundError(f"Файл не найден: {file_path}")
        
        if not path.suffix.lower() in ['.docx', '.doc']:
            logger.error(f"Неверный формат файла: {file_path}")
            raise ValueError(f"Ожидается DOCX файл, получен: {path.suffix}")
        
        logger.info(f"Парсинг DOCX: {file_path}")
        
        try:
            doc = Document(path)
            
            paragraphs = []
            full_text_parts = []
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)
                    full_text_parts.append(text)
            
            # Также извлекаем текст из таблиц
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            paragraphs.append(cell_text)
                            full_text_parts.append(cell_text)
            
            full_text = "\n\n".join(full_text_parts)
            
            if not full_text.strip():
                logger.warning(f"DOCX файл пуст: {file_path}")
            
            # Создаём DocumentText
            doc_text = DocumentText(
                content=full_text,
                paragraphs=paragraphs,
                source_file=path.name
            )
            
            # Извлекаем возраст и температуру из текста
            from nlp.context_analyzer import ContextAnalyzer
            
            analyzer = ContextAnalyzer(
                negation_patterns=[],
                exclusion_patterns=[],
                window_size=5
            )
            
            doc_text.age = analyzer.extract_age(full_text)
            doc_text.temperature = analyzer.extract_temperature(full_text)
            
            logger.info(
                f"DOCX распарсен: {len(paragraphs)} абзацев, "
                f"возраст={doc_text.age}, температура={doc_text.temperature}"
            )
            
            return doc_text
            
        except Exception as e:
            logger.error(f"Ошибка парсинга DOCX {file_path}: {e}")
            raise ValueError(f"Ошибка парсинга DOCX: {e}")
    
    def parse_bytes(self, data: bytes, filename: str = "document.docx") -> DocumentText:
        """Извлечение текста из DOCX данных в памяти.
        
        Args:
            data: Байты DOCX файла.
            filename: Имя файла (для логирования).
            
        Returns:
            DocumentText с извлечённым текстом.
        """
        import io
        
        logger.info(f"Парсинг DOCX из памяти: {filename}")
        
        try:
            doc = Document(io.BytesIO(data))
            
            paragraphs = []
            full_text_parts = []
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)
                    full_text_parts.append(text)
            
            # Извлекаем текст из таблиц
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            paragraphs.append(cell_text)
                            full_text_parts.append(cell_text)
            
            full_text = "\n\n".join(full_text_parts)
            
            doc_text = DocumentText(
                content=full_text,
                paragraphs=paragraphs,
                source_file=filename
            )
            
            # Извлекаем метаданные
            from nlp.context_analyzer import ContextAnalyzer
            
            analyzer = ContextAnalyzer([], [])
            doc_text.age = analyzer.extract_age(full_text)
            doc_text.temperature = analyzer.extract_temperature(full_text)
            
            return doc_text
            
        except Exception as e:
            logger.error(f"Ошибка парсинга DOCX из памяти: {e}")
            raise ValueError(f"Ошибка парсинга DOCX: {e}")
